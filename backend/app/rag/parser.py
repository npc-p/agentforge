"""文档解析：把不同格式统一抽取为纯文本。

只做「抽取」不做「理解」——结构化理解交给切分与检索层，
解析层保持单一职责，便于按格式独立扩展与测试。
"""

from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO
from typing import Any

from app.core.exceptions import AppError

SUPPORTED_EXT = {".txt", ".md", ".markdown", ".csv", ".log", ".json", ".pdf", ".docx"}
MAX_CHARS = 2_000_000


@dataclass(slots=True)
class ParsedDocument:
    text: str
    ext: str
    meta: dict[str, Any] = field(default_factory=dict)

    @property
    def char_count(self) -> int:
        return len(self.text)


def _decode(data: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gb18030", "big5", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _parse_pdf(data: bytes) -> tuple[str, dict[str, Any]]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise AppError("未安装 pypdf，无法解析 PDF", code="parser_missing") from exc

    reader = PdfReader(BytesIO(data))
    pages: list[str] = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:  # noqa: BLE001 - 个别页损坏不应中断整篇
            pages.append("")
    text = "\n\n".join(p.strip() for p in pages if p.strip())
    return text, {"pages": len(reader.pages), "pdf_meta": dict(reader.metadata or {})}


def _parse_docx(data: bytes) -> tuple[str, dict[str, Any]]:
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:  # pragma: no cover
        raise AppError("未安装 python-docx，无法解析 DOCX", code="parser_missing") from exc

    document = Document(BytesIO(data))
    blocks: list[str] = []
    table_rows = 0

    # 按 body 中的真实顺序遍历段落与表格。
    # 若分别读 document.paragraphs 与 document.tables，所有表格会被排到文末，
    # 从而脱离它所属的标题与上下文——切分后语义错位，检索与引用都会失真。
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = Paragraph(child, document).text.strip()
            if text:
                blocks.append(text)
        elif tag == "tbl":
            table = Table(child, document)
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    blocks.append(" | ".join(cells))
                    table_rows += 1

    return "\n".join(blocks), {"tables": len(document.tables), "table_rows": table_rows}


def extract_text(filename: str, data: bytes) -> ParsedDocument:
    suffix = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ".txt"
    if suffix not in SUPPORTED_EXT:
        raise AppError(
            f"不支持的文件类型: {suffix}",
            code="unsupported_file_type",
            detail={"supported": sorted(SUPPORTED_EXT)},
        )

    meta: dict[str, Any] = {"size_bytes": len(data)}
    if suffix == ".pdf":
        text, extra = _parse_pdf(data)
        meta.update(extra)
    elif suffix == ".docx":
        text, extra = _parse_docx(data)
        meta.update(extra)
    else:
        text = _decode(data)
        meta["lines"] = text.count("\n") + 1

    text = text.replace("\x00", "").strip()
    if not text:
        raise AppError("文档解析后内容为空，请确认文件是否为扫描件或加密文件", code="empty_document")
    if len(text) > MAX_CHARS:
        text = text[:MAX_CHARS]
        meta["truncated"] = True

    return ParsedDocument(text=text, ext=suffix.lstrip("."), meta=meta)
